"""Generate fictional six-format fixtures and a fixed evaluation question set."""

import argparse
import json
from pathlib import Path

import pymupdf
from docx import Document as WordDocument
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches


CASES = [
    {"id": "txt_1", "question": "ECHO-17 夜间支持工单需要在多久内首次响应？", "expected_source": "Operations.txt", "required_patterns": [r"45\s*分钟|四十五分钟"], "answerable": True},
    {"id": "txt_2", "question": "哪一类工单使用 ECHO-17 流程？", "expected_source": "Operations.txt", "required_patterns": ["夜间支持"], "answerable": True},
    {"id": "md_1", "question": "AURORA-52 的备份间隔和保留期限分别是多少？", "expected_source": "BackupPolicy.md", "required_patterns": [r"6\s*小时|六小时", r"21\s*天|二十一天"], "answerable": True},
    {"id": "docx_1", "question": "BRIDGE-04 新员工知识库培训持续几天？", "expected_source": "Onboarding.docx", "required_patterns": [r"3\s*天|三天"], "answerable": True},
    {"id": "docx_2", "question": "BRIDGE-04 结业时要完成几道模拟问答？", "expected_source": "Onboarding.docx", "required_patterns": [r"两道|2\s*道"], "answerable": True},
    {"id": "pdf_1", "question": "PULSE-63 的回滚时限是几分钟？", "expected_source": "ReleaseGuide.pdf", "required_patterns": [r"12\s*(分钟|minutes)"], "answerable": True},
    {"id": "pdf_2", "question": "PULSE-63 在什么条件下要求回滚？", "expected_source": "ReleaseGuide.pdf", "required_patterns": ["连续两次|两次连续|two consecutive"], "answerable": True},
    {"id": "pptx_1", "question": "MAPLE-81 产品演示安排在什么日期？", "expected_source": "DemoPlan.pptx", "required_patterns": [r"2026.{0,8}11.{0,8}20|November 20, 2026"], "answerable": True},
    {"id": "pptx_2", "question": "MAPLE-81 的演示交付物是什么？", "expected_source": "DemoPlan.pptx", "required_patterns": ["RAG", "看板"], "answerable": True},
    {"id": "xlsx_1", "question": "演示传感器的库存数量和仓库是什么？", "expected_source": "Inventory.xlsx", "required_patterns": ["37", "华东-02"], "answerable": True},
    {"id": "xlsx_2", "question": "演示路由器的库存数量是多少？", "expected_source": "Inventory.xlsx", "required_patterns": ["14"], "answerable": True},
    {"id": "no_answer_1", "question": "PULSE-63 的正式生产发布日期是什么？", "expected_source": None, "required_patterns": [], "answerable": False},
    {"id": "no_answer_2", "question": "AURORA-52 指定了哪一种加密算法？", "expected_source": None, "required_patterns": [], "answerable": False},
]


def generate(root: Path) -> list[Path]:
    if root.exists() and any(root.iterdir()):
        raise FileExistsError(f"输出目录不为空，请使用新目录：{root}")
    root.mkdir(parents=True, exist_ok=True)

    (root / "Operations.txt").write_text(
        "虚构演示资料：夜间支持工单使用流程代号 ECHO-17。\n"
        "值班团队收到夜间支持工单后，必须在 45 分钟内完成首次响应。\n"
        "首次响应是确认已收到工单，不等于解决故障。\n", encoding="utf-8"
    )
    (root / "BackupPolicy.md").write_text(
        "# 虚构备份制度 AURORA-52\n\n"
        "演练环境每 6 小时执行一次备份，备份文件保留 21 天。\n"
        "制度只描述备份频率与保留期限，没有指定任何加密算法。\n", encoding="utf-8"
    )

    word = WordDocument()
    word.add_heading("虚构入职培训 BRIDGE-04", 0)
    word.add_paragraph("新员工的知识库培训持续 3 天。")
    word.add_paragraph("结业时每人需要完成两道模拟问答，并在答案中指出资料来源。")
    word.save(root / "Onboarding.docx")

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text(
        (72, 72),
        "Fictional release guide PULSE-63.\n"
        "If two consecutive health checks fail, the operator must roll back within 12 minutes.\n"
        "The production launch date is not specified in this guide.", fontsize=11,
    )
    pdf.save(root / "ReleaseGuide.pdf")
    pdf.close()

    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4))
    box.text = "虚构展示计划 MAPLE-81\n产品演示日期：2026 年 11 月 20 日。\n演示交付物：RAG 检索看板和来源引用清单。"
    slides.save(root / "DemoPlan.pptx")

    book = Workbook()
    sheet = book.active
    sheet.title = "演示库存"
    sheet.append(["物品", "数量", "仓库"])
    sheet.append(["演示传感器", 37, "华东-02"])
    sheet.append(["演示路由器", 14, "华南-05"])
    book.save(root / "Inventory.xlsx")

    (root / "Cases.json").write_text(json.dumps(CASES, ensure_ascii=False, indent=2), encoding="utf-8")
    return [root / name for name in ("Operations.txt", "BackupPolicy.md", "Onboarding.docx", "ReleaseGuide.pdf", "DemoPlan.pptx", "Inventory.xlsx")]


def main():
    parser = argparse.ArgumentParser(description="生成六种格式的虚构评估资料")
    parser.add_argument("--output-dir", type=Path, default=Path("E:/Codex/RagEvaluationGenerated"))
    args = parser.parse_args()
    files = generate(args.output_dir)
    print(f"已生成 {len(files)} 份文件和 {len(CASES)} 道题：{args.output_dir}")


if __name__ == "__main__":
    main()
