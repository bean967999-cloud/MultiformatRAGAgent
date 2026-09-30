
from pathlib import Path

from pptx import Presentation
from langchain_core.documents import Document

from .base import BaseLoader

#PPT 本质上是由很多 shape文本框 组成的
class PPTLoader(BaseLoader):

    def load(self,file_path:Path) -> list[Document]:

        #打开这个 PPT 文件
        presentation = Presentation(file_path)

        documents = []

        #下面遍历两次，第一次是遍历每一个幻灯片，第二次是遍历每一个文本框，一个幻灯片可能会有很多的文本框
        #遍历每一页 PPT
        for slide_number,slide in enumerate(presentation.slides):     #遍历的时候，同时获取编号,遍历所有的幻灯片，同时给每个元素一个编号

            #当前 PPT 页面所有文本的列表
            texts = []
            #一页 PPT 里面可能有很多个文本框，slide.shapes 是当前 PPT 页面的所有形状/元素
            for shape in slide.shapes:
                #只处理文本框，是文本就执行下面的操作   判断这个 shape 有没有 text 属性
                if hasattr(shape,"text"):
                    #shape.text获取当前这个 PPT 元素里面的文字
                    #当前这个shape的文本
                    shape_text = shape.text.strip()
                    if shape_text:
                        texts.append(shape_text)

            slide_text = "\n".join(texts)
            #如果这个幻灯片有文本，就创建一个文档
            if slide_text:
                document = Document(
                    page_content=slide_text,
                    metadata = {
                        "source": str(file_path),
                        "file_type": "pptx",
                        "slide":slide_number+1
                    }
                )

                documents.append(document)

        return documents
