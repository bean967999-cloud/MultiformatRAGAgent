
from pathlib import Path

from openpyxl import load_workbook
from langchain_core.documents import Document

from .base import BaseLoader

class ExcelLoader(BaseLoader):

    def load(self,file_path:Path) -> list[Document]:

        #打开这个 Excel 文件
        workbook = load_workbook(
            file_path,
            data_only=True
        )

        documents = []

        for sheet_name in workbook.sheetnames:
            #获取当前这个工作表
            sheet = workbook[sheet_name]

            rows = list(sheet.values)

            if not rows:
                continue

            #第一行是表头，其他行是数据行
            headers = rows[0]

            for row_number,row in enumerate(rows[1:],start=2):

                content = []

                for key,value in zip(headers,row):

                    if key is not None and value is not None:
                        content.append(f"{key}:{value}")

                text = "\n".join(content)

                if text:
                    document = Document(
                        page_content=text,
                        metadata = {
                            "source": str(file_path),
                            "file_type": "xlsx",
                            "sheet":sheet_name,
                            "row":row_number
                        }
                    )

                    documents.append(document)


        return documents
