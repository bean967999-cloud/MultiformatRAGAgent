'''
Word 文件不是一个普通的文本文件，而是一个具有段落、表格、样式等结构的文档
现在是V1阶段，暂时不需要解析 Word 的标题、表格、代码块等复杂结构
"\n".join(列表)   用换行符，把列表中的字符串连接起来。
'''

from pathlib import Path
from docx import Document  as DocxDocument
from langchain_core.documents import Document
from .base import BaseLoader

class DOCXLoader(BaseLoader):
    #定义一个叫load的方法，它接收一个文件路径Path，最后返回一个Document列表
    def load(self,file_path:Path) ->list[Document]:     #返回值类型提示返回列表，往里面装Document对象
        doc = DocxDocument(file_path)   #打开指定的 Word 文件;这个 docx 对象里面就包含了 Word 文档的内容

        #准备一个容器，用来保存 Word 的所有段落文字
        paragraphs = []     #创建了一个空列表,准备把 Word 里面读取到的每一段文字放进去
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        text = "\n".join(paragraphs)

        document = Document(
            page_content=text,
            metadata = {
                'source':str(file_path),
                'file_type':'docx'
            }
        )

        return [document]
